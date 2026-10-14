"""
Photos of paper enquiry forms -> new rows in your existing tracking sheet.

Setup:
    pip install anthropic openpyxl pillow
    export ANTHROPIC_API_KEY="your-key"      (Windows: set ANTHROPIC_API_KEY=your-key)

Usage:
    python forms_to_excel.py photos/ tracker.xlsx
        -> reads the photos, writes a COPY called tracker_filled.xlsx (original untouched)

    python forms_to_excel.py photos/ tracker.xlsx --in-place    # write into the original
    python forms_to_excel.py photos/ tracker.xlsx --rotate 90   # if photos are sideways
    python forms_to_excel.py photos/ tracker.xlsx --dial-code 233   # default dial code (233 is the default)

How it works:
  * Each paper has two sides (student side, parent side). The script reads which side is which
    from the photo itself, so either order works. Sides are paired like this:
      - if EVERY filename starts with a number (1_front, 1_back, 2_front, ...), the number
        decides which photos belong together;
      - otherwise two neighbouring photos (in filename order) that are one student side and one
        parent side are paired.
  * The columns are found by reading your header row, and the dropdown lists (Gender,
    Education Qualification, Program of Interest) are read straight from the sheet's own data
    validation (or from dropdowns.json), so the script only writes values your dropdown accepts.
    Close matches are accepted, e.g. a program written without "BSc"/"BBA".
  * Cells it is unsure about are highlighted yellow with a comment naming the photo to check.
  * STATUS, REMARKS and COUNSELLOR are never touched.
  * A small log (<output>.log.csv) records which photos went into which row.
"""

import argparse
import base64
import csv
import difflib
import io
import json
import re
import time
from pathlib import Path

import anthropic
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from PIL import Image, ImageOps

MODEL = "claude-sonnet-5-5"  # change if your account uses a different model
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp"}
MAX_SIDE = 1800

# JSON key returned by the model -> internal column key
STUDENT_MAP = {
    "name": "student_name",
    "email": "student_email",
    "dial_code": "dial_code",
    "contact": "student_contact",
    "gender": "gender",
    "education_qualification": "education",
    "program_of_interest": "program",
}
PARENT_MAP = {
    "name": "parent_name",
    "contact": "parent_contact",
    "email": "parent_email",
    "occupation": "parent_occupation",
}
DROPDOWN_KEYS = ("gender", "education", "program")
YELLOW = PatternFill("solid", start_color="FFF2CC")


# --------------------------------------------------------------------------- sheet reading
def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().upper()


def find_columns(ws):
    """Locate the header row and map each field to its column number."""
    for header_row in range(1, 11):
        headers = {c.column: norm(c.value) for c in ws[header_row] if c.value}
        if "NAME OF STUDENT" in headers.values():
            break
    else:
        raise SystemExit("Could not find a header row containing 'NAME OF STUDENT'.")

    parent_col = next((c for c, h in headers.items() if h.startswith("NAME OF PARENT")), None)
    if parent_col is None:
        raise SystemExit("Could not find the 'NAME OF PARENT/GUARDIAN' column.")

    cols = {"parent_name": parent_col}
    for c, h in sorted(headers.items()):
        side = "parent" if c > parent_col else "student"
        if h == "NAME OF STUDENT":
            cols.setdefault("student_name", c)
        elif h == "EMAIL ADDRESS":
            cols.setdefault(f"{side}_email", c)
        elif h == "CONTACT":
            cols.setdefault(f"{side}_contact", c)
        elif h == "DIAL CODE":
            cols.setdefault("dial_code", c)
        elif h == "GENDER":
            cols.setdefault("gender", c)
        elif h.startswith("EDUCATION"):
            cols.setdefault("education", c)
        elif h.startswith("PROGRAM"):
            cols.setdefault("program", c)
        elif h == "OCCUPATION":
            cols.setdefault("parent_occupation", c)

    missing = [k for k in list(STUDENT_MAP.values()) + list(PARENT_MAP.values()) if k not in cols]
    if missing:
        raise SystemExit(f"Could not find these columns in the header row: {missing}")
    return header_row, cols


def dropdown_options(wb, ws, header_row, col):
    """Read the allowed values of a column's dropdown (data validation), if it has one."""
    probe = ws.cell(row=header_row + 1, column=col).coordinate
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list" or probe not in dv.sqref:
            continue
        formula = (dv.formula1 or "").strip()
        if formula.startswith('"'):  # inline list: "A,B,C"
            return [o.strip() for o in formula.strip('"').split(",") if o.strip()]
        try:  # a range such as Lists!$A$2:$A$20
            ref = formula.lstrip("=")
            sheet = ws
            if "!" in ref:
                sheet_name, ref = ref.rsplit("!", 1)
                sheet = wb[sheet_name.strip("'")]
            cells = sheet[ref.replace("$", "")]
            if not isinstance(cells, tuple):
                cells = ((cells,),)
            values = [c.value for row in cells for c in row]
            return [str(v).strip() for v in values if v not in (None, "")]
        except Exception:
            return []
    return []


def first_empty_row(ws, header_row, cols):
    r = header_row + 1
    while ws.cell(r, cols["student_name"]).value or ws.cell(r, cols["parent_name"]).value:
        r += 1
    return r


# --------------------------------------------------------------------------- the AI part
def build_prompt(options):
    def pick(key, what):
        opts = options.get(key)
        if opts:
            return f"{what}. Choose EXACTLY one of these values: {json.dumps(opts)}. " \
                   "Students often leave out prefixes such as BSc or BBA or spell loosely: pick " \
                   "the option that is closest in meaning to what is written. " \
                   "Only if it is genuinely ambiguous or matches nothing, return exactly what is " \
                   "written (the script will then leave that cell empty). Never invent a value " \
                   "that is not in the list."
        return f"{what}. Copy as written."

    return f"""You are reading a photo of a paper enquiry form that was filled in by hand.
The photo may be rotated or sideways; read it in whichever orientation makes the text readable.

There are two form types.

A) "Student Details" slip -> form_type "student" with these keys:
   name, email, dial_code, contact, gender, education_qualification, program_of_interest
   - gender: {pick("gender", "the ticked M / F box")}
   - education_qualification: {pick("education", "the ticked option(s) in the qualification section (SS1/SS2/SS3 under WASSCE/SSCE/NECO; English IB / French IB / American HS under IGCSE), or handwriting if written instead")}
     Hint: a tick on SS1, SS2 or SS3 means the WASSCE/SSCE/NECO row, so choose the dropdown
     option that contains WASSCE or NECO. A tick on English IB, French IB or American HS means
     that specific option, so choose the closest matching dropdown value.
   - program_of_interest: {pick("program", "the program written on the line")}
   - dial_code: the country code ONLY if one is written with the phone number (digits only, no +
     or 00). Otherwise "".
   - contact: the phone number digits only, WITHOUT the country code and WITHOUT a leading 0.
   (Ignore any other fields on the slip, such as school name, date of birth or address.)

B) "Parent/Guardian Details" slip -> form_type "parent" with these keys:
   name, contact, email, occupation
   - contact: digits only, WITHOUT a leading 0 and without country code.

Rules:
- Copy handwriting exactly. Do not fix spelling or guess missing letters.
- Blank field -> "".
- If you are unsure how to read a field, still give your best reading and add that key to
  "uncertain_fields". Emails with any ambiguous character must be marked uncertain.
- Emails in lowercase.

Return ONLY a JSON object (no markdown fences) with: form_type, uncertain_fields (list of key
names), and every key listed for that form type."""


def load_image_b64(path, rotate):
    img = ImageOps.exif_transpose(Image.open(path))
    if rotate:
        img = img.rotate(rotate, expand=True)  # counter-clockwise degrees
    img = img.convert("RGB")
    img.thumbnail((MAX_SIDE, MAX_SIDE))
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=90)
    return base64.standard_b64encode(buf.getvalue()).decode()


def extract(client, path, rotate, prompt):
    image_b64 = load_image_b64(path, rotate)
    resp, text = None, ""
    for _attempt in range(2):  # one automatic retry if the reply is not readable data
        resp = client.messages.create(
            model=MODEL,
            max_tokens=1000,
            messages=[{
                "role": "user",
                "content": [
                    {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg",
                                                  "data": image_b64}},
                    {"type": "text", "text": prompt},
                ],
            }],
        )
        text = "".join(b.text for b in resp.content if b.type == "text").strip()
        start, end = text.find("{"), text.rfind("}")  # tolerate text/fences around the JSON
        if start != -1 and end > start:
            try:
                return json.loads(text[start:end + 1])
            except json.JSONDecodeError:
                pass
    raise ValueError(f"no readable data in the reply (stop reason: {resp.stop_reason}; "
                     f"reply began: {text[:100]!r})")


# --------------------------------------------------------------------------- helpers
def phone_to_number(value):
    """'0740 759 345' -> 740759345 (matches how numbers appear in the sheet)."""
    digits = re.sub(r"\D", "", str(value or ""))
    digits = digits[1:] if digits.startswith("0") else digits
    return int(digits) if digits else None


_DEGREE_WORDS = {"bsc", "bba", "beng", "ba", "bs", "b", "sc", "bachelor", "of", "in", "degree"}


def _core_tokens(text):
    """Lower-case words of a program/qualification, ignoring degree prefixes and '&' vs 'and'."""
    t = str(text).casefold().replace("&", " and ")
    words = re.findall(r"[a-z0-9]+", t)
    core = [w for w in words if w not in _DEGREE_WORDS]
    return core or words


def snap_to_option(value, options):
    """Return (best matching dropdown option, ok_flag). Only clear matches are accepted."""
    if not options or not value:
        return value, True
    text = str(value).strip()
    for o in options:  # 1) exact, ignoring capitalization
        if o.casefold() == text.casefold():
            return o, True

    want = _core_tokens(text)
    cores = {o: _core_tokens(o) for o in options}

    # 2) same words once degree prefixes / '&' are ignored
    same = [o for o, c in cores.items() if c == want]
    if len(same) == 1:
        return same[0], True

    # 3) everything the student wrote appears in exactly one option
    #    (e.g. "Marketing" -> "BBA Marketing"; "WASSCE" -> "WASSCE/NECO")
    contained = [o for o, c in cores.items() if want and set(want) <= set(c)]
    if len(contained) == 1:
        return contained[0], True

    # 4) spelling differences ("A level" vs "A-LEVELS", "Entreprenuership"): clearly closest option
    want_s = " ".join(want)
    scored = sorted(((difflib.SequenceMatcher(None, want_s, " ".join(c)).ratio(), o)
                     for o, c in cores.items()), reverse=True)
    if scored and scored[0][0] >= 0.8 and (len(scored) == 1 or scored[0][0] - scored[1][0] >= 0.08):
        return scored[0][1], True
    return value, False


def resolve_sheet(wb, name):
    """Find the tab, ignoring capitalization and stray spaces; list the real tab names if not found."""
    if not name:
        return wb.active
    if name in wb.sheetnames:
        return wb[name]
    matches = [n for n in wb.sheetnames if n.strip().casefold() == name.strip().casefold()]
    if len(matches) == 1:
        return wb[matches[0]]
    raise SystemExit("Tab not found. The tabs in this workbook are:\n"
                     + "\n".join(f"  {n!r}" for n in wb.sheetnames)
                     + "\nCopy one exactly (the quotes show any spaces) and pass it with --sheet.")


def natural_key(path):
    """Sort like Windows Explorer's 'Name' sort: IMG_2 comes before IMG_10."""
    parts = re.split(r"(\d+)", path.name)
    return [int(t) if i % 2 else t.lower() for i, t in enumerate(parts)]


def taken_key(path):
    """Sort by the time the photo was taken (EXIF), else by file modified time."""
    try:
        exif = Image.open(path).getexif()
        raw = exif.get_ifd(0x8769).get(36867) or exif.get(306)  # DateTimeOriginal / DateTime
        if raw:
            return (str(raw), natural_key(path))
    except Exception:
        pass
    return (time.strftime("%Y:%m:%d %H:%M:%S", time.localtime(path.stat().st_mtime)), natural_key(path))


def pair_slips(items):
    """items: [(filename, data)] -> list of {'student', 'parent', 'student_file', 'parent_file'}."""
    def blank():
        return {"student": None, "student_file": "", "parent": None, "parent_file": ""}

    def put(rec, kind, fname, data):
        rec[kind], rec[f"{kind}_file"] = data, fname

    def kind_of(data):
        return data.get("form_type") if data.get("form_type") in ("student", "parent") else None

    # 1) every filename starts with a number: photos sharing that number belong together
    nums = [re.match(r"\s*(\d+)", fname) for fname, _ in items]
    if items and all(nums):
        groups = {}
        for (fname, data), m in zip(items, nums):
            groups.setdefault(int(m.group(1)), []).append((fname, data))
        rows = []
        for key in sorted(groups):
            rec = blank()
            for fname, data in groups[key]:
                k = kind_of(data)
                if k and rec[k] is None:
                    put(rec, k, fname, data)
                else:
                    print(f"SKIP {fname}: unknown form, or a second {k or 'slip'} with number {key}")
            if rec["student"] or rec["parent"]:
                rows.append(rec)
        return rows

    # 2) otherwise pair neighbours that are one student side and one parent side (either order)
    rows, i = [], 0
    while i < len(items):
        fname, data = items[i]
        k = kind_of(data)
        if k is None:
            print(f"SKIP {fname}: could not tell which form this is")
            i += 1
            continue
        rec = blank()
        put(rec, k, fname, data)
        if i + 1 < len(items):
            f2, d2 = items[i + 1]
            k2 = kind_of(d2)
            if k2 and k2 != k:
                put(rec, k2, f2, d2)
                i += 1
        rows.append(rec)
        i += 1
    return rows


# --------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("photos", type=Path, help="folder with the form photos")
    ap.add_argument("workbook", type=Path, help="your tracking sheet (.xlsx)")
    ap.add_argument("-o", "--output", type=Path, help="where to save (default: <name>_filled.xlsx)")
    ap.add_argument("--in-place", action="store_true", help="write straight into the original file")
    ap.add_argument("--sheet", help="sheet name (default: the first/active sheet)")
    ap.add_argument("--rotate", type=int, default=0, help="rotate photos CCW by this many degrees")
    ap.add_argument("--dial-code", default="233", help="dial code to use when the form has none (default 233; pass \"\" to leave blank)")
    ap.add_argument("--sort", choices=["name", "date"], default="name",
                    help="photo order: by filename (default, same as Explorer Name sort) or by time taken")
    ap.add_argument("--options-file", type=Path,
                    help="JSON file with dropdown values, used when the sheet has no dropdowns "
                         "(default: dropdowns.json in the current folder, if it exists)")
    args = ap.parse_args()

    out = args.workbook if args.in_place else (
        args.output or args.workbook.with_name(args.workbook.stem + "_filled.xlsx"))

    wb = load_workbook(args.workbook)
    ws = resolve_sheet(wb, args.sheet)
    header_row, cols = find_columns(ws)

    options = {k: dropdown_options(wb, ws, header_row, cols[k]) for k in DROPDOWN_KEYS}
    source = {k: "sheet" for k in DROPDOWN_KEYS}

    # A downloaded copy of an online sheet often loses its dropdowns. In that case use the
    # values from dropdowns.json (you create it yourself; it stays on your computer).
    opt_path = args.options_file or Path("dropdowns.json")
    file_opts = {}
    if opt_path.exists():
        file_opts = json.loads(opt_path.read_text(encoding="utf-8"))
    for k in DROPDOWN_KEYS:
        if not options[k] and file_opts.get(k):
            options[k] = [str(v).strip() for v in file_opts[k] if str(v).strip()]
            source[k] = opt_path.name

    for k in DROPDOWN_KEYS:
        print(f"Dropdown '{k}': {len(options[k])} options (from {source[k]})" if options[k]
              else f"Dropdown '{k}': NO options known - text will be copied as written")

    photos = sorted((p for p in args.photos.iterdir() if p.suffix.lower() in IMAGE_EXTS),
                    key=taken_key if args.sort == "date" else natural_key)
    if not photos:
        raise SystemExit(f"No images found in {args.photos}")

    client = anthropic.Anthropic()  # reads ANTHROPIC_API_KEY from the environment
    prompt = build_prompt(options)

    items = []
    for p in photos:
        try:
            items.append((p.name, extract(client, p, args.rotate, prompt)))
            print(f"read  {p.name}")
        except Exception as e:  # keep going if one photo fails
            print(f"FAIL  {p.name}: {e}")

    r = first_empty_row(ws, header_row, cols)
    log = []
    for rec in pair_slips(items):
        for kind, mapping in (("student", STUDENT_MAP), ("parent", PARENT_MAP)):
            data, fname = rec[kind], rec[f"{kind}_file"]
            if not data:
                continue
            uncertain = set(data.get("uncertain_fields", []))
            for key, colkey in mapping.items():
                value = data.get(key, "")
                written = value
                if colkey in DROPDOWN_KEYS:
                    value, ok = snap_to_option(value, options.get(colkey))
                    if not ok:
                        # Not one of the sheet's allowed dropdown values: leave the cell empty,
                        # but flag it and note what the form said so you can decide by hand.
                        cell = ws.cell(r, cols[colkey])
                        cell.fill = YELLOW
                        cell.comment = Comment(
                            f"Not in dropdown. Form says: {value} (photo: {fname})", "forms_to_excel")
                        continue
                if colkey in ("student_contact", "parent_contact"):
                    value = phone_to_number(value)
                elif colkey == "dial_code":
                    digits = re.sub(r"\D", "", str(value or ""))
                    if not digits and data.get("contact"):  # form had no code: use the default
                        digits = re.sub(r"\D", "", str(args.dial_code or ""))
                    value = int(digits) if digits else None
                if value in (None, ""):
                    continue
                cell = ws.cell(r, cols[colkey], value)
                if colkey in ("student_contact", "parent_contact", "dial_code"):
                    cell.number_format = "0"  # show full digits, never 5.38E+08
                if colkey in DROPDOWN_KEYS and str(written).strip().casefold() != str(value).casefold():
                    cell.comment = Comment(f"Form says: {written} (matched to closest option; "
                                           f"photo: {fname})", "forms_to_excel")
                if key in uncertain:
                    cell.fill = YELLOW
                    cell.comment = Comment(f"Check against photo: {fname}", "forms_to_excel")
        if rec["student"] is None:
            ws.cell(r, cols["student_name"]).comment = Comment("Parent slip with no student slip", "forms_to_excel")
            ws.cell(r, cols["student_name"]).fill = YELLOW
        elif rec["parent"] is None:
            ws.cell(r, cols["parent_name"]).comment = Comment("No parent slip found for this student", "forms_to_excel")
            ws.cell(r, cols["parent_name"]).fill = YELLOW
        log.append([r, rec["student_file"], rec["parent_file"]])
        r += 1

    try:
        wb.save(out)
    except PermissionError:
        backup = out.with_name(out.stem + "_new" + out.suffix)
        wb.save(backup)
        print(f"Could not write to {out} (is it open in Excel or another program?).")
        print(f"Your results were saved to {backup} instead. Close the original and copy rows over.")
        out = backup
    log_path = out.with_suffix(".log.csv")
    with open(log_path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["sheet_row", "student_photo", "parent_photo"])
        w.writerows(log)
    print(f"Added {len(log)} row(s). Saved {out} (log: {log_path})")
    print("Check every yellow cell against its photo before relying on the data.")


if __name__ == "__main__":
    main()
